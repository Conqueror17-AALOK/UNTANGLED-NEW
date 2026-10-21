import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import landscape, A3
from reportlab.pdfgen import canvas
from reportlab.lib import colors

def draw_poly(c, points, fill=0, stroke=1):
    p = c.beginPath()
    p.moveTo(points[0][0], points[0][1])
    for pt in points[1:]:
        p.lineTo(pt[0], pt[1])
    p.close()
    c.drawPath(p, fill=fill, stroke=stroke)

def draw_title_block(c, width, height, title, dwg_no, sheet_str):
    c.setStrokeColor(colors.black)
    c.setLineWidth(2)
    c.rect(30, 30, width - 60, height - 60)
    
    tb_w, tb_h = 360, 90
    tb_x = width - 30 - tb_w
    tb_y = 30
    c.rect(tb_x, tb_y, tb_w, tb_h)
    c.line(tb_x, tb_y + 45, tb_x + tb_w, tb_y + 45)
    c.line(tb_x + 220, tb_y, tb_x + 220, tb_y + 45)
    
    c.setFont("Helvetica-Bold", 14)
    c.drawString(tb_x + 15, tb_y + 60, "UNTANGLE DEMO PROCESS")
    c.setFont("Helvetica", 10)
    c.drawString(tb_x + 15, tb_y + 48, title)
    
    c.setFont("Helvetica-Bold", 10)
    c.drawString(tb_x + 15, tb_y + 25, f"DWG: {dwg_no}")
    c.setFont("Helvetica", 9)
    c.drawString(tb_x + 15, tb_y + 10, "REV: 01   DATE: 2026-10-21")
    c.drawString(tb_x + 230, tb_y + 25, f"SHEET: {sheet_str}")
    c.drawString(tb_x + 230, tb_y + 10, "SCALE: NONE")

def create_sample_pid_1(filename):
    c = canvas.Canvas(filename, pagesize=landscape(A3))
    w, h = landscape(A3)
    
    draw_title_block(c, w, h, "FEED PRETREATMENT & PUMPING SYSTEM", "DWG-UNT-001", "1 OF 1")
    
    c.setStrokeColor(colors.black)
    c.setFillColor(colors.black)
    
    # 1. Tank TK-101
    tk_x, tk_y, tk_w, tk_h = 100, 300, 140, 280
    c.setLineWidth(2.5)
    c.rect(tk_x, tk_y, tk_w, tk_h)
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(tk_x + tk_w/2, tk_y + tk_h/2 + 20, "TK-101")
    c.setFont("Helvetica", 10)
    c.drawCentredString(tk_x + tk_w/2, tk_y + tk_h/2, "FEED STORAGE")
    c.drawCentredString(tk_x + tk_w/2, tk_y + tk_h/2 - 15, "15,000 L")
    
    # 2. Feed line into TK-101
    c.setLineWidth(2)
    c.line(40, tk_y + tk_h - 40, tk_x, tk_y + tk_h - 40)
    c.drawString(45, tk_y + tk_h - 30, "RAW CRUDE INLET")
    
    # 3. Headspace Pressure Transmitter (PT-101)
    c.line(tk_x + tk_w/2, tk_y + tk_h, tk_x + tk_w/2, tk_y + tk_h + 50)
    c.circle(tk_x + tk_w/2, tk_y + tk_h + 70, 20)
    c.line(tk_x + tk_w/2 - 20, tk_y + tk_h + 70, tk_x + tk_w/2 + 20, tk_y + tk_h + 70)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(tk_x + tk_w/2, tk_y + tk_h + 75, "PT")
    c.drawCentredString(tk_x + tk_w/2, tk_y + tk_h + 58, "101")
    
    # PSV-101 on tank top
    c.line(tk_x + 25, tk_y + tk_h, tk_x + 25, tk_y + tk_h + 40)
    draw_poly(c, [(tk_x + 15, tk_y + tk_h + 40), (tk_x + 35, tk_y + tk_h + 40), (tk_x + 25, tk_y + tk_h + 55)], fill=0)
    c.setFont("Helvetica", 8)
    c.drawString(tk_x + 30, tk_y + tk_h + 45, "PSV-101")
    
    # LT-101 on Tank side
    c.line(tk_x - 30, tk_y + 50, tk_x, tk_y + 50)
    c.line(tk_x - 30, tk_y + 200, tk_x, tk_y + 200)
    c.line(tk_x - 30, tk_y + 50, tk_x - 30, tk_y + 200)
    c.circle(tk_x - 60, tk_y + 125, 20)
    c.line(tk_x - 80, tk_y + 125, tk_x - 40, tk_y + 125)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(tk_x - 60, tk_y + 130, "LT")
    c.drawCentredString(tk_x - 60, tk_y + 115, "101")
    c.line(tk_x - 40, tk_y + 125, tk_x - 30, tk_y + 125)
    
    # 4. Discharge Process Line from TK-101 to Pumps
    suction_y = tk_y + 30
    c.line(tk_x + tk_w, suction_y, 420, suction_y)
    
    # XV-101 Solenoid Isolation Valve
    draw_poly(c, [(260, suction_y - 12), (284, suction_y + 12), (284, suction_y - 12), (260, suction_y + 12)], fill=0)
    c.rect(266, suction_y + 12, 12, 12)
    c.setFont("Helvetica", 8)
    c.drawString(258, suction_y - 25, "XV-101")
    
    # Strainer FL-101
    draw_poly(c, [(340, suction_y - 15), (370, suction_y), (340, suction_y + 15)], fill=0)
    c.setFont("Helvetica", 8)
    c.drawString(335, suction_y - 25, "FL-101")
    
    # Split to Pump A (P-101A) and Standby Pump B (P-101B)
    c.line(420, suction_y, 420, suction_y + 80)
    c.line(420, suction_y, 420, suction_y - 80)
    
    # Pump A
    p_a_x, p_a_y = 520, suction_y + 80
    c.line(420, p_a_y, p_a_x, p_a_y)
    c.circle(p_a_x + 25, p_a_y, 25)
    draw_poly(c, [(p_a_x + 25, p_a_y + 25), (p_a_x + 45, p_a_y + 25), (p_a_x + 25, p_a_y + 10)], fill=0)
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(p_a_x + 25, p_a_y - 38, "P-101A")
    c.setFont("Helvetica", 8)
    c.drawCentredString(p_a_x + 25, p_a_y - 48, "DUTY PUMP")
    
    # Pump B
    p_b_x, p_b_y = 520, suction_y - 80
    c.line(420, p_b_y, p_b_x, p_b_y)
    c.circle(p_b_x + 25, p_b_y, 25)
    draw_poly(c, [(p_b_x + 25, p_b_y + 25), (p_b_x + 45, p_b_y + 25), (p_b_x + 25, p_b_y + 10)], fill=0)
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(p_b_x + 25, p_b_y - 38, "P-101B")
    c.setFont("Helvetica", 8)
    c.drawCentredString(p_b_x + 25, p_b_y - 48, "STANDBY PUMP")
    
    # Check valves CK-101A & CK-101B
    c.line(p_a_x + 25, p_a_y + 25, p_a_x + 25, p_a_y + 60)
    c.line(p_b_x + 25, p_b_y + 25, p_b_x + 25, p_b_y + 60)
    
    # Common discharge header
    header_y = p_a_y + 60
    c.line(p_b_x + 25, p_b_y + 60, 620, p_b_y + 60)
    c.line(620, p_b_y + 60, 620, header_y)
    c.line(p_a_x + 25, header_y, 750, header_y)
    
    # Heat Exchanger E-101
    hx_x, hx_y, hx_w, hx_h = 750, header_y - 30, 80, 60
    c.roundRect(hx_x, hx_y, hx_w, hx_h, 10)
    c.line(hx_x + 15, hx_y, hx_x + 15, hx_y + hx_h)
    c.line(hx_x + hx_w - 15, hx_y, hx_x + hx_w - 15, hx_y + hx_h)
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(hx_x + hx_w/2, hx_y + hx_h/2 - 4, "E-101")
    c.setFont("Helvetica", 7)
    c.drawCentredString(hx_x + hx_w/2, hx_y - 12, "PRE-HEATER")
    
    # Flow transmitter FT-101
    c.line(hx_x + hx_w, header_y, 900, header_y)
    c.line(870, header_y, 870, header_y + 40)
    c.circle(870, header_y + 60, 20)
    c.line(850, header_y + 60, 890, header_y + 60)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(870, header_y + 65, "FT")
    c.drawCentredString(870, header_y + 48, "101")
    
    # Flow Indicating Controller FIC-101
    c.setDash(3, 3)
    c.line(870, header_y + 80, 870, header_y + 130)
    c.setDash()
    c.circle(870, header_y + 150, 20)
    c.setDash(2, 2)
    c.line(850, header_y + 150, 890, header_y + 150)
    c.setDash()
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(870, header_y + 155, "FIC")
    c.drawCentredString(870, header_y + 138, "101")
    
    # Control Valve CV-101
    c.line(900, header_y, 980, header_y)
    cv_x = 940
    draw_poly(c, [(cv_x - 12, header_y - 10), (cv_x + 12, header_y + 10), (cv_x + 12, header_y - 10), (cv_x - 12, header_y + 10)], fill=0)
    c.line(cv_x, header_y, cv_x, header_y + 20)
    c.arc(cv_x - 15, header_y + 20, cv_x + 15, header_y + 40, 0, 180)
    c.line(cv_x - 15, header_y + 20, cv_x + 15, header_y + 20)
    c.setFont("Helvetica", 8)
    c.drawCentredString(cv_x, header_y - 22, "CV-101")
    
    # Signal line from FIC-101 to CV-101
    c.setDash(3, 3)
    c.line(890, header_y + 150, cv_x, header_y + 150)
    c.line(cv_x, header_y + 150, cv_x, header_y + 30)
    c.setDash()
    
    # Temperature Transmitter TT-101
    c.line(980, header_y, 1050, header_y)
    c.line(1005, header_y, 1005, header_y - 35)
    c.circle(1005, header_y - 55, 20)
    c.line(985, header_y - 55, 1025, header_y - 55)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(1005, header_y - 50, "TT")
    c.drawCentredString(1005, header_y - 67, "101")
    
    # Manual Hand Valve HV-77
    hv_x = 1040
    draw_poly(c, [(hv_x - 10, header_y - 8), (hv_x + 10, header_y + 8), (hv_x + 10, header_y - 8), (hv_x - 10, header_y + 8)], fill=0)
    c.line(hv_x, header_y, hv_x, header_y + 15)
    c.line(hv_x - 8, header_y + 15, hv_x + 8, header_y + 15)
    c.setFont("Helvetica", 8)
    c.drawCentredString(hv_x, header_y - 20, "HV-77")
    
    # Off-page connector to Unit 200
    c.line(1050, header_y, 1100, header_y)
    draw_poly(c, [(1100, header_y - 15), (1135, header_y - 15), (1155, header_y), (1135, header_y + 15), (1100, header_y + 15)], fill=0)
    c.setFont("Helvetica-Bold", 8)
    c.drawCentredString(1125, header_y - 4, "TO UNIT 200")
    c.setFont("Helvetica", 7)
    c.drawCentredString(1125, header_y - 25, "DWG-UNT-002")
    
    c.showPage()
    c.save()
    print(f"Generated {filename}")

def create_sample_pid_multipage(filename):
    c = canvas.Canvas(filename, pagesize=landscape(A3))
    w, h = landscape(A3)
    
    # Page 1: Pre-treatment
    draw_title_block(c, w, h, "MULTISHEET SYSTEM - FEED PREPARATION", "DWG-UNT-002", "1 OF 2")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, 700, "SHEET 1: CRUDE FEED PRETREATMENT")
    
    # Tank V-201
    c.rect(150, 320, 140, 240)
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(220, 450, "V-201")
    c.setFont("Helvetica", 9)
    c.drawCentredString(220, 430, "FLASH SEPARATOR")
    
    # Line to off-page connector
    c.line(290, 440, 950, 440)
    draw_poly(c, [(950, 425), (1010, 425), (1035, 440), (1010, 455), (950, 455)], fill=0)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(985, 436, "TO SHEET 2")
    c.setFont("Helvetica", 8)
    c.drawCentredString(985, 410, "LINE 10-PR-201")
    
    c.showPage()
    
    # Page 2: Reaction Loop
    draw_title_block(c, w, h, "MULTISHEET SYSTEM - REACTION TRAIN", "DWG-UNT-002", "2 OF 2")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(100, 700, "SHEET 2: HYDROTREATER & REACTION TRAIN")
    
    # From sheet 1 off-page connector
    draw_poly(c, [(100, 440), (125, 425), (185, 425), (185, 455), (125, 455)], fill=0)
    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(150, 436, "FROM SHEET 1")
    c.setFont("Helvetica", 8)
    c.drawCentredString(150, 410, "LINE 10-PR-201")
    
    # Line to Reactor R-201
    c.line(185, 440, 450, 440)
    c.rect(450, 280, 160, 320)
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(530, 450, "R-201")
    c.setFont("Helvetica", 10)
    c.drawCentredString(530, 430, "HYDROTREATER")
    
    c.showPage()
    c.save()
    print(f"Generated {filename}")

def create_sample_io_list(filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IO_List"
    
    headers = ["Tag", "Description", "Signal Type", "Engineering Units", "Range Low", "Range High", "P&ID Ref", "Location"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    rows = [
        ["TK-101", "Feed Storage Tank", "Equipment", "Level %", 0, 100, "DWG-UNT-001", "Tank Farm Area 1"],
        ["P-101A", "Crude Feed Duty Pump", "415V MCC Motor", "RPM", 0, 2950, "DWG-UNT-001", "Pump House 1"],
        ["P-101B", "Crude Feed Standby Pump", "415V MCC Motor", "RPM", 0, 2950, "DWG-UNT-001", "Pump House 1"],
        ["PT-101", "Tank Headspace Pressure Transmitter", "4-20 mA AI", "barg", 0, 5.0, "DWG-UNT-001", "Tank Farm Area 1"],
        ["FT-101", "Discharge Main Flow Transmitter", "4-20 mA AI", "m3/h", 0, 120.0, "DWG-UNT-001", "Battery Limit Unit 100"],
        ["FIC-101", "Flow Indicating Controller", "DCS Loop", "m3/h", 0, 120.0, "DWG-UNT-001", "DCS Rack Room"],
        ["LT-101", "Feed Tank Level Transmitter", "4-20 mA AI", "%", 0, 100.0, "DWG-UNT-001", "Tank Farm Area 1"],
        ["TT-101", "Effluent Temperature Transmitter", "RTD AI", "deg C", -20, 250.0, "DWG-UNT-001", "Battery Limit Unit 100"],
        ["CV-101", "Feed Flow Pneumatic Control Valve", "4-20 mA AO", "% Open", 0, 100.0, "DWG-UNT-001", "Piping Rack Unit 100"],
        ["XV-101", "Suction Emergency Isolation Solenoid Valve", "24V DC DO", "Open/Close", 0, 1, "DWG-UNT-001", "Tank Farm Area 1"],
        ["PSV-101", "Vessel Safety Relief Valve", "Mechanical", "barg", 4.5, 4.5, "DWG-UNT-001", "Tank Headspace"],
        ["FT-310", "Secondary Recycle Flow Transmitter (Missing from Dwg)", "4-20 mA AI", "m3/h", 0, 50.0, "DWG-UNT-001", "Unit 300 Recycle"]
    ]
    
    for r in rows:
        ws.append(r)
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    wb.save(filename)
    print(f"Generated {filename}")

if __name__ == "__main__":
    create_sample_pid_1("contracts/samples/sample_pid_01.pdf")
    create_sample_pid_multipage("contracts/samples/sample_pid_02_multipage.pdf")
    create_sample_io_list("contracts/samples/sample_io_list.xlsx")
